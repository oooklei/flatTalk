# -*- coding: utf-8 -*-
"""将广西养老政策文件分类转换为 Markdown（政策原文/解读 → 养老政策；表格 → 办事指引）。"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

try:
    import docx
except ImportError:
    print("请先安装: pip install python-docx")
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[2]  # 广西养老政策文件及办事指引
SRC = ROOT / "广西养老政策文件"
POLICY_OUT = Path(__file__).resolve().parent / "原文"
FORM_OUT = ROOT / "知识库" / "办事指引" / "表格说明"
REPORT = Path(__file__).resolve().parent / "_convert_report.json"

FORM_KEYWORDS = (
    "申请表",
    "告知书",
    "汇总表",
    "承诺书",
    "机构要求",
    "审核表",
)


def classify(stem: str) -> str:
    if "政策解读" in stem:
        return "解读"
    if any(k in stem for k in FORM_KEYWORDS):
        return "表格"
    return "原文"


def extract_doc_no(text: str, stem: str) -> str:
    for src in (stem, text[:800]):
        m = re.search(r"(桂民[发规函]〔\d{4}〕\d+号)", src)
        if m:
            return m.group(1)
    return ""


def docx_to_text(path: Path) -> str:
    d = docx.Document(str(path))
    parts: list[str] = []
    for p in d.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
            # 去重相邻重复单元格（合并单元格常见）
            cleaned: list[str] = []
            for c in cells:
                if not cleaned or cleaned[-1] != c:
                    cleaned.append(c)
            if any(cleaned):
                parts.append(" | ".join(cleaned))
    return "\n\n".join(parts)


def try_doc_via_word(path: Path) -> str | None:
    """尝试用 Word COM 读取 .doc；失败返回 None。"""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return None
    word = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(path.resolve()))
        text = doc.Content.Text
        doc.Close(False)
        # Word 用 \r 换行
        lines = [ln.strip() for ln in text.replace("\r", "\n").split("\n") if ln.strip()]
        return "\n\n".join(lines)
    except Exception as e:
        print("word-com fail", path.name, e)
        return None
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass


def safe_name(name: str, max_len: int = 90) -> str:
    name = re.sub(r'[\\/:*?"<>|]', "_", name)
    name = re.sub(r"\s+", " ", name).strip()
    if len(name) > max_len:
        name = name[: max_len - 1].rstrip() + "…"
    return name


def write_md(out_path: Path, title: str, kind: str, doc_no: str, rel: str, body: str) -> None:
    meta = [
        f"# {title}",
        "",
        "## 元数据",
        "",
        f"| 项目 | 内容 |",
        f"|------|------|",
        f"| 类型 | {kind} |",
        f"| 文号 | {doc_no or '（文件中未单独标注或见正文）'} |",
        f"| 来源文件 | `{rel}` |",
        "",
        "## 正文",
        "",
        body.strip(),
        "",
    ]
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(meta), encoding="utf-8")


def main() -> int:
    if not SRC.exists():
        print("missing source", SRC)
        return 1
    POLICY_OUT.mkdir(parents=True, exist_ok=True)
    FORM_OUT.mkdir(parents=True, exist_ok=True)

    report = {
        "policy": [],
        "form": [],
        "skipped": [],
        "errors": [],
    }

    files = sorted(list(SRC.rglob("*.docx")) + list(SRC.rglob("*.doc")))
    # 排除临时文件
    files = [p for p in files if not p.name.startswith("~$")]

    for path in files:
        rel = str(path.relative_to(SRC))
        stem = path.stem
        kind = classify(stem)
        try:
            if path.suffix.lower() == ".docx":
                text = docx_to_text(path)
            else:
                text = try_doc_via_word(path) or ""
                if not text:
                    report["skipped"].append({"file": rel, "reason": ".doc 无法解析（需 Word/COM）"})
                    continue
        except Exception as e:
            report["errors"].append({"file": rel, "error": str(e)})
            print("fail", rel, e)
            continue

        if len(text.strip()) < 20:
            report["skipped"].append({"file": rel, "reason": "内容过短"})
            continue

        doc_no = extract_doc_no(text, stem) or extract_doc_no(rel, rel)
        # 同名表格（不同政策包）加父目录文号/前缀避免覆盖
        if kind == "表格" and path.parent != SRC:
            parent_tag = extract_doc_no(path.parent.name, path.parent.name) or safe_name(path.parent.name, 24)
            out_name = safe_name(f"{parent_tag}_{stem}") + ".md"
        else:
            out_name = safe_name(stem) + ".md"

        if kind == "表格":
            out_path = FORM_OUT / out_name
            write_md(out_path, stem, "表格说明", doc_no, rel, text)
            report["form"].append({"file": rel, "out": str(out_path.relative_to(ROOT)), "doc_no": doc_no})
        else:
            label = "政策解读" if kind == "解读" else "政策原文"
            out_path = POLICY_OUT / out_name
            write_md(out_path, stem, label, doc_no, rel, text)
            report["policy"].append(
                {"file": rel, "out": str(out_path.relative_to(ROOT)), "kind": label, "doc_no": doc_no}
            )
        print("ok", kind, out_name)

    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        "done policy=",
        len(report["policy"]),
        "form=",
        len(report["form"]),
        "skipped=",
        len(report["skipped"]),
        "errors=",
        len(report["errors"]),
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
