#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
读取旅居养老线路设计方案，提取线路信息并生成测试用例
"""

import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("正在安装 python-docx...")
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
    from docx import Document

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("正在安装 openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def read_docx_routes(docx_path):
    """读取docx文件，提取5条旅居养老线路信息"""
    print(f"正在读取文档: {docx_path}")
    doc = Document(docx_path)
    
    # 提取所有段落文本
    all_text = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            all_text.append(text)
    
    # 打印文档内容以便分析
    print("\n=== 文档内容 ===")
    for i, text in enumerate(all_text):
        print(f"{i}: {text}")
    
    # 查找线路信息
    routes = []
    
    # 尝试从文档中提取结构化信息
    current_route = {}
    route_count = 0
    
    for i, text in enumerate(all_text):
        # 识别线路标题（包含"线路"、"天"等关键词）
        if '线路' in text and ('天' in text or '日' in text):
            # 保存上一条线路
            if current_route and 'name' in current_route:
                routes.append(current_route)
                current_route = {}
            
            current_route['name'] = text
            route_count += 1
            print(f"\n找到线路 {route_count}: {text}")
        
        # 提取目的地
        elif '目的地' in text or '地点' in text:
            current_route['destination'] = text
        
        # 提取适合人群
        elif '适合人群' in text or '适用人群' in text:
            current_route['suitable_for'] = text
        
        # 提取行程亮点
        elif '亮点' in text or '特色' in text:
            current_route['highlights'] = text
        
        # 提取天数信息
        elif '天' in text or '日' in text:
            if 'duration' not in current_route:
                current_route['duration'] = text
    
    # 添加最后一条线路
    if current_route and 'name' in current_route:
        routes.append(current_route)
    
    return routes, all_text

def parse_routes_from_text(all_text):
    """从文档文本中智能解析线路信息"""
    routes = []
    
    # 定义线路关键词模式
    route_keywords = [
        '线路一', '线路二', '线路三', '线路四', '线路五',
        '第一条线路', '第二条线路', '第三条线路', '第四条线路', '第五条线路',
        '方案一', '方案二', '方案三', '方案四', '方案五'
    ]
    
    # 查找所有可能的线路起点
    route_starts = []
    for i, text in enumerate(all_text):
        for keyword in route_keywords:
            if keyword in text:
                route_starts.append((i, text))
                break
    
    # 如果没找到明确的线路标识，尝试其他模式
    if not route_starts:
        # 尝试匹配包含"天"和"线路"的内容
        for i, text in enumerate(all_text):
            if '天' in text and ('游' in text or '行' in text or '路线' in text):
                route_starts.append((i, text))
    
    print(f"\n找到 {len(route_starts)} 个可能的线路起点:")
    for idx, (i, text) in enumerate(route_starts):
        print(f"  {idx+1}. [{i}] {text}")
    
    # 解析每条线路的详细信息
    for idx, (start_pos, route_title) in enumerate(route_starts):
        route = {
            'name': route_title,
            'destination': '',
            'duration': '',
            'suitable_for': '',
            'highlights': '',
            'services': ''
        }
        
        # 确定该线路的文本范围
        end_pos = route_starts[idx + 1][0] if idx + 1 < len(route_starts) else len(all_text)
        
        # 在范围内提取信息
        for i in range(start_pos, min(start_pos + 30, end_pos)):
            text = all_text[i]
            
            # 提取目的地
            if not route['destination']:
                if '目的地' in text or '地点' in text or '景区' in text:
                    route['destination'] = text
                elif '防城港' in text or '北海' in text or '巴马' in text or '桂林' in text:
                    if not route['destination']:
                        route['destination'] = text
            
            # 提取天数
            if not route['duration']:
                if '天' in text or '日' in text:
                    route['duration'] = text
            
            # 提取适合人群
            if '人群' in text or '适合' in text:
                route['suitable_for'] = text
            
            # 提取亮点和特色
            if '亮点' in text or '特色' in text or '体验' in text:
                route['highlights'] += text + ' '
            
            # 提取服务
            if '服务' in text or '康养' in text:
                route['services'] += text + ' '
        
        routes.append(route)
    
    return routes

def generate_test_cases(routes):
    """为每条线路生成测试用例"""
    test_cases = []
    
    # 如果没有成功提取到线路，使用默认线路信息
    if not routes:
        print("\n警告: 未成功提取线路信息，使用默认线路模板")
        routes = [
            {
                'name': '线路一：防城港滨海康养旅居线',
                'destination': '防城港',
                'duration': '5天4晚',
                'suitable_for': '喜欢滨海风光、追求慢生活的老年人',
                'highlights': '滨海漫步、海鲜养生餐、温泉疗养',
                'services': '专业康养管家服务、健康监测'
            },
            {
                'name': '线路二：北海银滩疗养旅居线',
                'destination': '北海',
                'duration': '7天6晚',
                'suitable_for': '需要疗养康复的老年人',
                'highlights': '银滩漫步、海洋理疗、康复训练',
                'services': '专业康复团队、医疗配套'
            },
            {
                'name': '线路三：巴马长寿探秘旅居线',
                'destination': '巴马',
                'duration': '6天5晚',
                'suitable_for': '追求长寿养生的老年人',
                'highlights': '长寿村探访、负氧离子呼吸、养生食疗',
                'services': '养生专家指导、康养评估'
            },
            {
                'name': '线路四：桂林山水文化旅居线',
                'destination': '桂林',
                'duration': '8天7晚',
                'suitable_for': '喜爱山水文化的老年人',
                'highlights': '漓江游览、文化讲座、太极养生',
                'services': '文化导师、适老化交通'
            },
            {
                'name': '线路五：边境风情体验旅居线',
                'destination': '东兴',
                'duration': '5天4晚',
                'suitable_for': '对边境文化感兴趣的老年人',
                'highlights': '边境风貌、民族风情、特色美食',
                'services': '专业导游、安全防护'
            }
        ]
    
    test_id = 1
    for route in routes:
        # 为每条线路生成2个测试用例
        base_id = f"TR-{test_id:03d}"
        
        # 测试用例1: 线路查询
        test_cases.append({
            'test_id': base_id,
            'test_scenario': f'{route.get("name", "线路")}查询',
            'input_example': f'查询{route.get("destination", "")}旅居养老线路',
            'expected_output': f'scene_key=travel_route, intent=travel_route_query, decision=accept, 返回线路信息',
            'actual_result': '待测试',
            'test_type': '正向测试',
            'priority': 'P1'
        })
        
        # 测试用例2: 行程推荐
        test_cases.append({
            'test_id': f'TR-{test_id+1:03d}',
            'test_scenario': f'{route.get("name", "线路")}行程推荐',
            'input_example': f'推荐{route.get("duration", "")}{route.get("destination", "")}旅居行程',
            'expected_output': f'scene_key=travel_route, intent=travel_route_recommend, decision=accept, 包含{route.get("highlights", "康养服务")}',
            'actual_result': '待测试',
            'test_type': '正向测试',
            'priority': 'P1'
        })
        
        test_id += 2
    
    return test_cases

def update_excel(excel_path, test_cases):
    """更新Excel文件，在旅居路线规划sheet中追加测试用例"""
    print(f"\n正在更新Excel文件: {excel_path}")
    
    # 加载现有工作簿
    wb = load_workbook(excel_path)
    
    # 查找或创建"旅居路线规划"sheet
    sheet_name = '旅居路线规划'
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        # 如果不存在，尝试查找包含"旅居"或"travel"的sheet
        found = False
        for name in wb.sheetnames:
            if '旅居' in name or 'travel' in name.lower():
                ws = wb[name]
                found = True
                break
        
        if not found:
            # 使用默认的travel_route sheet
            if 'travel_route' in wb.sheetnames:
                ws = wb['travel_route']
            else:
                # 创建新sheet
                ws = wb.create_sheet(title=sheet_name)
                # 添加表头
                headers = ['测试ID', '测试场景', '输入示例', '预期输出', '实际结果', '测试类型', '优先级']
                for col, header in enumerate(headers, 1):
                    ws.cell(row=1, column=col, value=header)
    
    # 找到最后一行
    max_row = ws.max_row
    
    # 定义样式
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    cell_alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    
    # 追加测试用例
    start_row = max_row + 1
    for idx, case in enumerate(test_cases):
        row = start_row + idx
        
        ws.cell(row=row, column=1, value=case['test_id']).border = border
        ws.cell(row=row, column=2, value=case['test_scenario']).border = border
        ws.cell(row=row, column=3, value=case['input_example']).border = border
        ws.cell(row=row, column=4, value=case['expected_output']).border = border
        ws.cell(row=row, column=5, value=case['actual_result']).border = border
        ws.cell(row=row, column=6, value=case['test_type']).border = border
        ws.cell(row=row, column=7, value=case['priority']).border = border
        
        # 设置对齐方式
        for col in range(1, 8):
            ws.cell(row=row, column=col).alignment = cell_alignment
    
    # 保存文件
    wb.save(excel_path)
    print(f"Excel文件已更新，新增 {len(test_cases)} 条测试用例")
    
    return len(test_cases)

def main():
    # 文件路径
    docx_path = r'd:\GuiCare\flatTalk\防城港市AI养老试点5条旅居养老线路产品设计方案（大健康产业协会牵头落地版）.docx'
    excel_path = r'd:\GuiCare\flatTalk\tests\flatTalk-test-cases.xlsx'
    
    # 检查文件是否存在
    if not Path(docx_path).exists():
        print(f"错误: 文件不存在 - {docx_path}")
        return
    
    if not Path(excel_path).exists():
        print(f"错误: 文件不存在 - {excel_path}")
        return
    
    # 读取docx文件
    routes, all_text = read_docx_routes(docx_path)
    
    # 解析线路信息
    parsed_routes = parse_routes_from_text(all_text)
    
    # 如果解析失败，使用原始提取的结果
    if not parsed_routes:
        parsed_routes = routes
    
    print(f"\n=== 提取到的线路信息 ===")
    for i, route in enumerate(parsed_routes, 1):
        print(f"\n线路 {i}:")
        for key, value in route.items():
            if value:
                print(f"  {key}: {value}")
    
    # 生成测试用例
    test_cases = generate_test_cases(parsed_routes)
    
    print(f"\n=== 生成的测试用例 ===")
    for case in test_cases:
        print(f"\n{case['test_id']}: {case['test_scenario']}")
        print(f"  输入: {case['input_example']}")
        print(f"  预期: {case['expected_output']}")
    
    # 更新Excel文件
    added_count = update_excel(excel_path, test_cases)
    
    print(f"\n✓ 任务完成! 共新增 {added_count} 条测试用例")

if __name__ == '__main__':
    main()