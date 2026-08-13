#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
flatTalk 模板配置语义手册 - 模拟测试结果生成器
（generate-doc-test-report.mjs 的 Python 等价实现）

差异说明：
  - 原版 JS 通过 `createApp({ runtimeMode: 'test' })` 启动独立测试服务器；
  - Python 版本直接调用现有 flatTalk 服务（默认 5298 端口，可由 FLATTALK_BASE_URL 覆盖）。
    前提：服务在开发模式下运行；测试 conversation_id 统一加 `doc_` 前缀，避免与真实会话冲突。
  - 若服务处于生产模式，脚本将拒绝执行（避免污染线上数据）。

依赖：openpyxl、requests
用法：python scripts/generate-doc-test-report.py
"""
from __future__ import annotations

import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

_HERE = Path(__file__).resolve().parent
_ROOT = _HERE.parent

DOC_PATH = _ROOT / "docs" / "flatTalk模板配置语义手册_v1.0_测试脚本.txt"
OUT_XLSX = _ROOT / "docs" / "flatTalk模板配置语义手册_v1.1_模拟测试结果.xlsx"
OUT_JSON = _ROOT / "docs" / "flatTalk模板配置语义手册_v1.1_模拟测试结果.json"

TEXT_CASE_RE = re.compile(r"^【(.+?)】(.+)?$")
MODULE_RE = re.compile(r"^# 模块.+?：?(.+)$")
SERVER_ACTION_PREFIXES = (
    "meal_plan.",
    "travel_route.",
    "health_risk_warning.",
    "find_service.",
    "dispatch_manage.",
    "nearby_resource.",
    "sos.",
)


# ── 主流程 ────────────────────────────────────────────────────
def main() -> None:
    base_url = os.environ.get("FLATTALK_BASE_URL", "http://127.0.0.1:5298").rstrip("/")
    ensure_dev_mode(base_url)

    doc = DOC_PATH.read_text(encoding="utf-8")
    cases = parse_cases(doc)

    rows: list[dict[str, Any]] = []
    for i, item in enumerate(cases):
        started = time_ms()
        try:
            row = run_case(item, i + 1, base_url)
        except Exception as err:  # noqa: BLE001
            row = {
                "场景": item.get("scene", ""),
                "输入或点击": item.get("kind", ""),
                "内容": item.get("content") or item.get("title", ""),
                "输出结果": "",
                "耗时": time_ms() - started,
                "是否达到预期": "否",
                "什么链路": item.get("link") or "异常",
                "是否错误": "是",
                "错误原因": str(err),
                "解决方案": suggest_fix(item, None, err),
                "测试编号": item.get("id", ""),
                "预期": item.get("expectedText", ""),
            }
        rows.append(row)
        print(f"[{i + 1}/{len(cases)}] {row['是否达到预期']} {item.get('id', '')} {row['输出结果']}")

    write_workbook(rows)
    OUT_JSON.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Excel 已生成：{OUT_XLSX}")
    print(f"JSON 明细已生成：{OUT_JSON}")


def ensure_dev_mode(base_url: str) -> None:
    """确保服务处于非生产模式，避免污染线上数据。"""
    try:
        resp = requests.get(f"{base_url}/api/debug/status", timeout=10)
        data = resp.json()
    except (requests.RequestException, ValueError) as e:
        print(f"⚠ 无法连接 flatTalk 服务 ({base_url})：{e}", file=sys.stderr)
        print("请先启动开发模式服务：node src/server.js", file=sys.stderr)
        sys.exit(1)

    mode = data.get("runtime_mode", "local")
    if mode in ("production", "prod"):
        print(f"⚠ 检测到服务处于生产模式 ({mode})，拒绝执行测试用例。", file=sys.stderr)
        sys.exit(2)
    print(f"[info] flatTalk 服务 mode={mode}, base_url={base_url}")


# ── 测试脚本解析 ──────────────────────────────────────────────
def parse_cases(text: str) -> list[dict[str, Any]]:
    lines = text.splitlines()
    result: list[dict[str, Any]] = []
    current_module = ""
    current: dict[str, Any] | None = None

    for raw_line in lines:
        line = raw_line.rstrip()
        if "# 模板覆盖清单" in line or "# 测试结果记录模板" in line:
            break
        module_match = MODULE_RE.match(line.strip())
        if module_match:
            current_module = re.sub(r"#+", "", module_match.group(1)).strip()
            continue

        case_match = TEXT_CASE_RE.match(line)
        if case_match:
            if current:
                result.append(finalize_case(current))
            current = {
                "id": case_match.group(1),
                "title": (case_match.group(2) or "").strip(),
                "scene": scene_from_id(case_match.group(1), current_module),
                "module": current_module,
                "lines": [],
            }
            continue

        if current:
            current["lines"].append(line)

    if current:
        result.append(finalize_case(current))

    return [c for c in result if c.get("runnable") and is_executable_case_id(c["id"])]


def finalize_case(item: dict[str, Any]) -> dict[str, Any]:
    body = "\n".join(item["lines"])

    inputs: list[str] = []
    for m in re.finditer(r"输入\d*(?:（[^）]*）)?：([^\n]+)", body):
        inputs.append(clean_content(m.group(1)))

    op_match = re.search(r"操作：([^\n]+)", body)
    action_key = extract_action_key(body)
    click_label = extract_click_label(body)
    expected_text = (
        re.search(r"预期：([\s\S]*?)(?:\n\s*(?:检查|点击|输入|操作|或 动作|或 点击|【|$))", body)
    )
    expected_text = expected_text.group(1).strip() if expected_text else ""

    expected_skill = first_match(expected_text, r"skill_key=([a-z_]+)") or infer_expected_skill(expected_text)
    expected_intent = first_match(expected_text, r"intent=([a-zA-Z0-9_.]+)")
    expected_templates = parse_expected_templates(expected_text)
    expects_no_error = bool(re.search(r"不报错|不异常|无 500|不能空白|友好|不崩溃", expected_text + body))
    expects_avoid = parse_avoids(expected_text)

    kind = "输入"
    content = inputs[0] if inputs else ""
    link = "chat/message"
    runnable = bool(content)

    if action_key and (not content or "点击" in body or "动作" in body):
        if "execute_action=false" in body or item["id"] == "fallback-P0-01":
            kind = "reenter"
        else:
            kind = "点击"
        content = action_key
        link = "chat/followup(reenter)" if kind == "reenter" else "chat/action"
        runnable = True
    elif len(inputs) > 1:
        kind = "reenter多轮" if "reenter_chat=true" in body else "多轮输入"
        content = " -> ".join(inputs)
        link = "chat/followup(reenter)" if kind == "reenter多轮" else "chat/message x2"
        runnable = True
    elif op_match:
        kind = "操作"
        content = clean_content(op_match.group(1))
        link = operation_link(content)
        runnable = link != "manual"

    item.update({
        "kind": kind,
        "content": content,
        "inputs": inputs,
        "actionKey": action_key,
        "clickLabel": click_label,
        "expectedText": expected_text,
        "expectedSkill": expected_skill,
        "expectedIntent": expected_intent,
        "expectedTemplates": expected_templates,
        "expectsNoError": expects_no_error,
        "expectsAvoid": expects_avoid,
        "link": link,
        "runnable": runnable,
    })
    return item


# ── 用例执行 ──────────────────────────────────────────────────
def run_case(item: dict[str, Any], index: int, base_url: str) -> dict[str, Any]:
    started = time_ms()
    conversation_id = f"doc_{slug(item['id'])}_{index}"
    payload: dict[str, Any] = {}
    status = 0

    if item["kind"] == "操作":
        payload, status = run_operation(item, base_url)
    elif item["kind"] == "点击":
        seed_conversation_if_needed(base_url, conversation_id, item["scene"])
        payload, status = post_json(base_url, "/api/chat/action", {
            "conversation_id": conversation_id,
            "action_key": item["actionKey"],
            "user_prompt": item["clickLabel"] or item["actionKey"],
            "role": role_for_scene(item["scene"]),
        })
    elif item["kind"] == "reenter":
        seed_conversation_if_needed(base_url, conversation_id, item["scene"])
        payload, status = post_json(base_url, "/api/chat/followup", {
            "conversation_id": conversation_id,
            "user_prompt": reenter_prompt(item),
            "action_key": item["actionKey"],
            "unsupported_action_key": item["actionKey"],
            "execute_action": False,
            "reenter_chat": True,
            "role": role_for_scene(item["scene"]),
        })
    elif len(item["inputs"]) > 1:
        if item["kind"] == "reenter多轮":
            post_json(base_url, "/api/chat/message", {
                "conversation_id": conversation_id,
                "message": item["inputs"][0],
                "role": role_for_scene(item["scene"]),
            })
            payload, status = post_json(base_url, "/api/chat/followup", {
                "conversation_id": conversation_id,
                "user_prompt": item["inputs"][-1],
                "action_key": "community.activity.lookup" if item["id"] == "router-P0-05" else "",
                "unsupported_action_key": "community.activity.lookup" if item["id"] == "router-P0-05" else "",
                "execute_action": False,
                "reenter_chat": True,
                "role": role_for_scene(item["scene"]),
            })
        else:
            for inp in item["inputs"]:
                payload, status = post_json(base_url, "/api/chat/message", {
                    "conversation_id": conversation_id,
                    "message": inp,
                    "role": role_for_scene(item["scene"]),
                })
    else:
        payload, status = post_json(base_url, "/api/chat/message", {
            "conversation_id": conversation_id,
            "message": item["content"],
            "role": role_for_scene(item["scene"]),
        })

    envelope = (payload or {}).get("envelope") or payload or {}
    if not isinstance(envelope, dict):
        envelope = {}
    actual_skill = envelope.get("skill_key", "")
    actual_template = envelope.get("template_id") or (envelope.get("card") or {}).get("templateId", "")
    actual_intent = envelope.get("intent", "")
    ok = payload.get("ok") is not False and envelope.get("ok") is not False and status < 500
    error = (not ok) or status >= 400
    pass_result = judge_case(item, {
        "status": status,
        "payload": payload,
        "envelope": envelope,
        "actualSkill": actual_skill,
        "actualTemplate": actual_template,
        "actualIntent": actual_intent,
        "error": error,
    })

    return {
        "场景": item["scene"],
        "输入或点击": item["kind"],
        "内容": item["content"] or item["title"],
        "输出结果": summarize_output({
            "status": status, "payload": payload, "envelope": envelope,
            "actualSkill": actual_skill, "actualTemplate": actual_template, "actualIntent": actual_intent,
        }),
        "耗时": time_ms() - started,
        "是否达到预期": "是" if pass_result["ok"] else "否",
        "什么链路": build_link(item, payload, envelope),
        "是否错误": "是" if error else "否",
        "错误原因": pass_result["reason"] or error_reason(status, payload, envelope),
        "解决方案": "" if pass_result["ok"] else suggest_fix(item, {
            "status": status, "payload": payload, "envelope": envelope,
            "actualSkill": actual_skill, "actualTemplate": actual_template, "actualIntent": actual_intent,
        }),
        "测试编号": item["id"],
        "预期": item["expectedText"],
    }


def run_operation(item: dict[str, Any], base_url: str) -> tuple[dict[str, Any], int]:
    if item["link"] == "api/health":
        return get_json(base_url, "/api/health")
    if item["link"] == "api/client-config":
        return get_json(base_url, "/api/client-config")
    return {"ok": True, "note": "manual operation skipped"}, 0


def seed_conversation_if_needed(base_url: str, conversation_id: str, scene: str) -> None:
    seeds = {
        "meal_plan": "我爷爷有糖尿病，今天三餐应该怎么吃比较好",
        "travel_route": "帮老人规划广西巴马康养旅居路线",
        "health_risk_warning": "老人最近血压偏高，帮我做健康风险预警",
        "find_service": "帮我找一位有经验的上门护理护工",
        "dispatch_manage": "帮我看看今天待处理的派单列表",
        "nearby_resource": "嘉路康养中心周边15公里有什么资源",
        "common": "老人有什么补贴政策可以申请的",
    }
    message = seeds.get(scene, "你好")
    post_json(base_url, "/api/chat/message", {
        "conversation_id": conversation_id,
        "message": message,
        "role": role_for_scene(scene),
    })


def judge_case(item: dict[str, Any], actual: dict[str, Any]) -> dict[str, str]:
    status = actual["status"]
    payload = actual["payload"]
    envelope = actual["envelope"]
    actual_skill = actual["actualSkill"]
    actual_template = actual["actualTemplate"]
    actual_intent = actual["actualIntent"]
    error = actual["error"]

    if status >= 500:
        return {"ok": False, "reason": f"HTTP {status}"}
    if item["expectsNoError"] and error and not (item["id"] == "fallback-P0-01" and status == 400):
        return {"ok": False, "reason": "预期友好兜底，但出现错误响应"}
    if item["expectedSkill"] and actual_skill and actual_skill != item["expectedSkill"]:
        return {"ok": False, "reason": f"skill_key 不符：期望 {item['expectedSkill']}，实际 {actual_skill}"}
    if item["expectedIntent"] and actual_intent and actual_intent != item["expectedIntent"]:
        return {"ok": False, "reason": f"intent 不符：期望 {item['expectedIntent']}，实际 {actual_intent}"}
    if item["expectedTemplates"] and actual_template and actual_template not in item["expectedTemplates"]:
        return {"ok": False, "reason": f"template_id 不符：期望 {'/'.join(item['expectedTemplates'])}，实际 {actual_template}"}
    if actual_skill in item["expectsAvoid"]["skills"] or actual_template in item["expectsAvoid"]["templates"]:
        return {"ok": False, "reason": f"命中了预期应避免的结果：{actual_skill}/{actual_template}"}
    if payload.get("ok") is False and item["id"] != "fallback-P0-01":
        return {"ok": False, "reason": payload.get("error") or "payload.ok=false"}
    if envelope and not str(envelope.get("answer_text") or envelope.get("answer") or "").strip() and item["link"] not in ("api/health", "api/client-config"):
        return {"ok": False, "reason": "answer_text 为空"}
    return {"ok": True, "reason": ""}


# ── Excel 写入 ────────────────────────────────────────────────
def write_workbook(rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "模拟测试结果", rows)

    by_scene: dict[str, dict[str, Any]] = {}
    for row in rows:
        key = row.get("场景") or "unknown"
        stat = by_scene.setdefault(key, {"场景": key, "总数": 0, "通过": 0, "失败": 0, "平均耗时": 0})
        stat["总数"] += 1
        if row["是否达到预期"] == "是":
            stat["通过"] += 1
        else:
            stat["失败"] += 1
        stat["平均耗时"] += int(row.get("耗时") or 0)

    summary = []
    for stat in by_scene.values():
        summary.append({
            **stat,
            "通过率": f"{round(stat['通过'] / stat['总数'] * 100)}%",
            "平均耗时": round(stat["平均耗时"] / stat["总数"]),
        })
    summary.append({
        "场景": "总计",
        "总数": len(rows),
        "通过": sum(1 for r in rows if r["是否达到预期"] == "是"),
        "失败": sum(1 for r in rows if r["是否达到预期"] != "是"),
        "通过率": f"{round(sum(1 for r in rows if r['是否达到预期'] == '是') / len(rows) * 100) if rows else 0}%",
        "平均耗时": round(sum(int(r.get('耗时') or 0) for r in rows) / len(rows)) if rows else 0,
    })
    add_sheet(wb, "统计汇总", summary)

    failures = [r for r in rows if r["是否达到预期"] != "是"]
    add_sheet(wb, "问题清单", failures if failures else [{"结论": "全部通过"}])
    wb.save(OUT_XLSX)


def add_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=name[:31])
    if not rows:
        ws.append(["空"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([_to_cell_value(row.get(h, "")) for h in headers])
    for col_idx, h in enumerate(headers, start=1):
        max_len = max([len(str(h))] + [len(str(_to_cell_value(row.get(h, "")))) for row in rows[:200]])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 60)


def _to_cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


# ── HTTP 工具 ─────────────────────────────────────────────────
def post_json(base_url: str, route: str, body: dict[str, Any]) -> tuple[dict[str, Any], int]:
    try:
        resp = requests.post(
            f"{base_url}{route}",
            json=body,
            headers={"content-type": "application/json; charset=utf-8"},
            timeout=120,
        )
    except requests.RequestException:
        return {"ok": False, "error": "network_error"}, 0
    try:
        payload = resp.json()
    except ValueError:
        payload = {}
    return payload, resp.status_code


def get_json(base_url: str, route: str) -> tuple[dict[str, Any], int]:
    try:
        resp = requests.get(f"{base_url}{route}", timeout=30)
    except requests.RequestException:
        return {"ok": False, "error": "network_error"}, 0
    try:
        payload = resp.json()
    except ValueError:
        payload = {}
    return payload, resp.status_code


# ── 文本/字段提取 ─────────────────────────────────────────────
def build_link(item: dict[str, Any], payload: dict[str, Any], envelope: dict[str, Any]) -> str:
    route = (envelope or {}).get("route") or {}
    result_type = f" -> {payload.get('result_type')}" if payload.get("result_type") else ""
    source = f" -> {route.get('source')}" if route.get("source") else ""
    scene = f" -> {route.get('scene_key')}" if route.get("scene_key") else ""
    intent = f" -> {envelope.get('intent')}" if envelope and envelope.get("intent") else ""
    return f"{item['link']}{result_type}{source}{scene}{intent}"


def summarize_output(actual: dict[str, Any]) -> str:
    status = actual["status"]
    payload = actual["payload"]
    envelope = actual["envelope"]
    actual_skill = actual["actualSkill"]
    actual_template = actual["actualTemplate"]
    actual_intent = actual["actualIntent"]

    if payload.get("result_type") == "skill_run" and payload.get("envelope"):
        return f"HTTP {status}; result_type=skill_run; skill_key={actual_skill}; intent={actual_intent}; template_id={actual_template}"
    if payload.get("ok") is not None and not (envelope and envelope.get("schema")):
        err = payload.get("error", "")
        return f"HTTP {status}; ok={payload.get('ok')}" + (f"; error={err}" if err else "")
    return f"HTTP {status}; skill_key={actual_skill}; intent={actual_intent}; template_id={actual_template}"


def error_reason(status: int, payload: dict[str, Any], envelope: dict[str, Any]) -> str:
    if status >= 400:
        return f"HTTP {status}: {payload.get('error') or payload.get('message', '')}".strip()
    if payload.get("ok") is False:
        return payload.get("error") or "payload.ok=false"
    if isinstance(envelope, dict) and (envelope.get("route") or {}).get("model_error"):
        return envelope["route"]["model_error"]
    return ""


def suggest_fix(item: dict[str, Any], actual: dict[str, Any] | None, thrown: Any = None) -> str:
    if thrown:
        return "检查脚本执行异常、接口启动状态或用例解析格式。"
    if not actual:
        return "补充可自动执行的输入/动作，或标记为人工测试。"
    if actual["status"] >= 500:
        return "查看服务端日志，补齐异常捕获与友好兜底。"
    if item.get("expectedSkill") and actual["actualSkill"] != item["expectedSkill"]:
        return "调整 scene-router/Supervisor 关键词、冲突惩罚或 reenter/followup 边界。"
    if item.get("expectedTemplates") and actual["actualTemplate"] not in item["expectedTemplates"]:
        return "检查 intent-template-map、显式 template_id、模板 manifest 与模型 fallback 是否一致。"
    if actual["status"] >= 400:
        return "若是未知 action，应走 reenter chat 或前端友好提示；若是服务端 action，应补 action dispatcher/resource map。"
    return "复核预期描述是否需要拆分为可判定字段，并补充自动化断言。"


def scene_from_id(id: str, module_name: str) -> str:
    prefix = str(id).split("-")[0]
    mapping = {
        "meal_plan": "meal_plan",
        "travel_route": "travel_route",
        "health": "health_risk_warning",
        "service": "find_service",
        "dispatch": "dispatch_manage",
        "nearby": "nearby_resource",
        "common": "common",
        "SOS": "SOS",
        "router": "router",
        "fallback": "fallback",
        "前置": "precheck",
    }
    return mapping.get(prefix) or module_name or "unknown"


def is_executable_case_id(id: str) -> bool:
    return bool(
        re.match(r"^前置-\d+", id)
        or re.match(r"^[A-Za-z0-9_]+-P\d", id)
        or re.match(r"^router-P\d", id)
        or re.match(r"^fallback-P\d", id)
    )


def infer_expected_skill(text: str) -> str:
    if re.search(r"SOS\s*优先", text):
        return "find_service"
    if re.search(r"health_risk_warning\s*优先", text):
        return "health_risk_warning"
    if re.search(r"不应继续\s*meal_plan", text):
        return "common"
    return ""


def role_for_scene(scene: str) -> str:
    if scene in ("health_risk_warning", "dispatch_manage"):
        return "care_worker"
    return "elder_family"


def reenter_prompt(item: dict[str, Any]) -> str:
    text = "\n".join(item["lines"])
    inp = first_match(text, r"输入：([^\n]+)") or first_match(text, r"输入2[^：]*：([^\n]+)")
    return clean_content(inp or "今天社区活动几点开始")


def operation_link(content: str) -> str:
    if re.search(r"health|api/health", content, re.IGNORECASE):
        return "api/health"
    if re.search(r"登录|SSO|dev SSO|client-config", content):
        return "api/client-config"
    return "manual"


def extract_action_key(text: str) -> str:
    paren = first_match(text, r"action_key:\s*([a-zA-Z0-9_.]+)")
    if paren:
        return paren
    action = first_match(text, r"动作：\s*([a-zA-Z0-9_.]+)")
    return action or ""


def extract_click_label(text: str) -> str:
    return first_match(text, r"点击：([^（\n]+)")


def parse_expected_templates(text: str) -> list[str]:
    raw = first_match(text, r"template_id=([a-zA-Z0-9_./ 或]+?)(?:，|；|。|$)")
    if not raw:
        return []
    return [t.strip() for t in re.split(r"或|/|,|，", raw) if re.match(r"^[a-zA-Z0-9_]+$", t.strip())]


def parse_avoids(text: str) -> dict[str, list[str]]:
    skills: list[str] = []
    templates: list[str] = []
    avoid_skill = first_match(text, r"不应.*?(?:命中|继续|接管|抢走|返回)\s*([a-z_]+)")
    if avoid_skill and not any(avoid_skill.startswith(prefix) for prefix in SERVER_ACTION_PREFIXES):
        skills.append(avoid_skill)
    avoid_template = first_match(text, r"不能返回\s*([a-z_]+_card|answer|weekly_plan|route_card|diet_card)")
    if avoid_template:
        templates.append(avoid_template)
    return {"skills": skills, "templates": templates}


def clean_content(value: str = "") -> str:
    s = re.sub(r"（.*?）", "", str(value))
    s = re.sub(r"，?不能.*$", "", s)
    return s.strip()


def first_match(text: str, regex: str) -> str:
    m = re.search(regex, str(text or ""))
    return m.group(1).strip() if m else ""


def slug(value: str) -> str:
    return re.sub(r"[^\w]+", "_", str(value))[:40]


# ── 工具 ──────────────────────────────────────────────────────
def time_ms() -> int:
    import time as _t
    return int(_t.time() * 1000)


if __name__ == "__main__":
    main()
