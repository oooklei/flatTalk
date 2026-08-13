#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
flatTalk 模板配置语义手册 - 自动审计报告生成器
（generate-template-audit-report.mjs 的 Python 等价实现）

依赖：
  - openpyxl：生成 .xlsx 工作簿
  - flattalk_client.FlatTalkClient：通过 HTTP 调用 Node.js 的 runLocalSkill / dispatchAction / getTraceLogger

用法：
  python scripts/generate-template-audit-report.py
环境要求：flatTalk 服务在本地 5298 端口（或 FLATTALK_BASE_URL 指定）以开发模式运行。
"""
from __future__ import annotations

import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# 让脚本既能从仓库根目录运行，也能从 scripts/ 目录运行
_HERE = Path(__file__).resolve().parent
_ROOT = _HERE.parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))

from flattalk_client import FlatTalkClient  # noqa: E402

# ── 常量 ──────────────────────────────────────────────────────
SCRIPT_PATH = _ROOT / "docs" / "flatTalk模板配置语义手册_v1.0_测试脚本.txt"
OUT_DIR = _ROOT / "docs" / "reports"
RUN_ID = f"template-audit-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}"


# ── 主流程 ────────────────────────────────────────────────────
def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    text = SCRIPT_PATH.read_text(encoding="utf-8")
    cases = parse_cases(text)
    started_at = datetime.now()

    client = FlatTalkClient()
    rows = run_cases(cases, client)
    ended_at = datetime.now()

    logs = collect_logs(rows, client)
    coverage = build_coverage(text, rows)
    summary = build_summary(rows, logs, started_at, ended_at)
    assessment = build_assessment(rows, logs, coverage)

    xlsx_path = OUT_DIR / f"{RUN_ID}.xlsx"
    json_path = OUT_DIR / f"{RUN_ID}.json"
    md_path = OUT_DIR / f"{RUN_ID}-评估意见.md"

    write_workbook(
        xlsx_path,
        summary=summary,
        rows=rows,
        logs=logs,
        coverage=coverage,
        assessment=assessment,
    )
    json_path.write_text(
        json.dumps(
            {
                "run_id": RUN_ID,
                "summary": summary,
                "rows": rows,
                "logs": logs,
                "coverage": coverage,
                "assessment": assessment,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    md_path.write_text(render_markdown(summary, assessment), encoding="utf-8")

    print(
        json.dumps(
            {
                "ok": True,
                "run_id": RUN_ID,
                "cases": len(rows),
                "xlsx": str(xlsx_path),
                "json": str(json_path),
                "markdown": str(md_path),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


# ── 测试脚本解析 ──────────────────────────────────────────────
def parse_cases(text: str) -> list[dict[str, Any]]:
    lines = text.splitlines()
    cases: list[dict[str, Any]] = []
    module_name = ""
    current: dict[str, Any] | None = None

    def flush() -> None:
        nonlocal current
        if not current:
            return
        normalize_case(current)
        cases.append(current)
        current = None

    for line in lines:
        module_match = re.match(r"^#\s*模块[^：:]*[：:]\s*(.+)$", line)
        if module_match:
            module_name = re.sub(r"#+", "", module_match.group(1)).strip()

        case_match = re.match(r"^【([^】]+)】(.+)$", line)
        if case_match:
            flush()
            current = {
                "id": case_match.group(1).strip(),
                "title": case_match.group(2).strip(),
                "module": module_name,
                "input": "",
                "action_key": "",
                "expected_skill": "",
                "expected_intent": "",
                "expected_template": "",
                "expected_output": "",
                "checks": [],
                "raw": [line],
            }
            continue
        if not current:
            continue
        current["raw"].append(line)
        trimmed = line.strip()
        if re.match(r"^输入[：:]", trimmed) and not current["input"]:
            current["input"] = re.sub(r"^输入[：:]\s*", "", trimmed).strip()
        if re.match(r"^(点击|动作)[：:]", trimmed) and not current["action_key"]:
            current["action_key"] = extract_action_key(trimmed)
        if re.match(r"^action_key\s*[:：]\s*([a-zA-Z0-9_.-]+)", trimmed) and not current["action_key"]:
            current["action_key"] = extract_action_key(trimmed)
        if re.match(r"^预期[：:]", trimmed):
            expected = re.sub(r"^预期[：:]\s*", "", trimmed).strip()
            current["expected_output"] = append_text(current["expected_output"], expected)
            if not current["expected_skill"]:
                current["expected_skill"] = pick_assign(expected, "skill_key")
            if not current["expected_intent"]:
                current["expected_intent"] = pick_assign(expected, "intent")
            if not current["expected_template"]:
                current["expected_template"] = pick_assign(expected, "template_id")
            if not current.get("expected_result_type"):
                current["expected_result_type"] = pick_assign(expected, "result_type") or ""
            current["negative_expectation"] = infer_negative_expectation(expected)
        if re.match(r"^检查[：:]", trimmed):
            current["checks"].append(re.sub(r"^检查[：:]\s*", "", trimmed).strip())
    flush()
    return [c for c in cases if c["input"] or c["action_key"] or c["expected_skill"] or c["expected_template"]]


def normalize_case(item: dict[str, Any]) -> None:
    if not item["action_key"]:
        item["action_key"] = extract_action_key("\n".join(item["raw"]))
    if not item["expected_template"] and item["expected_output"]:
        m = re.search(r"template_id\s*=\s*([a-zA-Z0-9_]+)", item["expected_output"])
        if m:
            item["expected_template"] = m.group(1)
    if not item["expected_skill"] and item["expected_output"]:
        m = re.search(r"skill_key\s*=\s*([a-zA-Z0-9_]+)", item["expected_output"])
        if m:
            item["expected_skill"] = m.group(1)
    item["kind"] = "chat" if item["input"] else "action"


def extract_action_key(text: str) -> str:
    m = re.search(r"action_key\s*[:：]\s*([a-zA-Z0-9_.-]+)", str(text or ""))
    return m.group(1) if m else ""


def pick_assign(text: str, key: str) -> str:
    m = re.search(rf"{key}\s*=\s*([a-zA-Z0-9_.-]+)", str(text or ""))
    return m.group(1) if m else ""


def infer_negative_expectation(text: str) -> str:
    if re.search(r"不应命中\s*nearby_resource", text):
        return "not_nearby_resource"
    if re.search(r"不应.*travel_route|不应强行当旅居规划", text):
        return "not_strong_travel_route"
    if re.search(r"不能返回\s*diet_card", text):
        return "not_diet_card"
    return ""


# ── 用例执行 ──────────────────────────────────────────────────
def run_cases(cases: list[dict[str, Any]], client: FlatTalkClient) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, item in enumerate(cases):
        conversation_id = re.sub(r"[^a-zA-Z0-9_.-]", "-", f"{RUN_ID}-{item['id']}")
        request_id = f"req-{RUN_ID}-{index + 1}"
        turn_id = f"turn-{RUN_ID}-{index + 1}"
        started = time.perf_counter()
        result: Any = None
        error = ""
        try:
            if item["kind"] == "action":
                result = client.dispatch_action(
                    action_key=item["action_key"],
                    params={},
                    conversation_id=conversation_id,
                    request_id=request_id,
                    turn_id=turn_id,
                )
                # 桥端点统一返回 { ok, result, source_status, ... }，真正的 dispatchAction 结果在 result 字段
                if isinstance(result, dict) and "result" in result and isinstance(result["result"], dict):
                    result = result["result"]
            else:
                context = {}
                if item["action_key"]:
                    context = {
                        "action_key": item["action_key"],
                        "action_params": {},
                        "followup_source": "script_inline_action",
                    }
                result = client.run_local_skill(
                    message=item["input"],
                    context=context,
                    conversation_id=conversation_id,
                    request_id=request_id,
                    turn_id=turn_id,
                )
                if isinstance(result, dict) and "result" in result and isinstance(result["result"], dict):
                    result = result["result"]
        except Exception as err:  # noqa: BLE001
            error = f"{type(err).__name__}: {err}"

        elapsed_ms = round((time.perf_counter() - started) * 1000)
        envelope = result.get("envelope", result) if isinstance(result, dict) else None
        if item["kind"] == "action" and isinstance(result, dict):
            envelope = result.get("envelope") or result
        actual = normalize_actual(item, result, envelope, elapsed_ms, error, conversation_id)
        rows.append({
            "序号": index + 1,
            "测试编号": item["id"],
            "模块": item["module"],
            "标题": item["title"],
            "类型": item["kind"],
            "输入或动作": item["input"] or item["action_key"],
            "action_key": item["action_key"],
            "预期skill_key": item["expected_skill"],
            "实际skill_key": actual["skill_key"],
            "预期intent": item["expected_intent"],
            "实际intent": actual["intent"],
            "预期template_id": item["expected_template"],
            "实际template_id": actual["template_id"],
            "预期输出检查": "；".join([item["expected_output"], *item["checks"]]).strip("；"),
            "实际输出摘要": actual["answer"],
            "路由是否正确": judge_route(item, actual),
            "模板是否正确": judge_template(item, actual),
            "输出是否合理": judge_output(item, actual),
            "动作是否同场景": judge_same_scene_actions(actual),
            "渲染是否正常": "通过" if actual["render_ok"] else "不通过",
            "result_type": actual["result_type"],
            "route_source": actual["route_source"],
            "route_decision": actual["route_decision"],
            "route_confidence": actual["route_confidence"],
            "model_used": actual["model_used"],
            "model_status": actual["model_status"],
            "knowledge_status": actual["knowledge_status"],
            "耗时ms": elapsed_ms,
            "conversation_id": conversation_id,
            "错误": error,
            "备注": build_remark(item, actual, error),
        })
    return rows


def normalize_actual(
    item: dict[str, Any],
    result: Any,
    envelope: dict[str, Any] | None,
    elapsed_ms: int,
    error: str,
    conversation_id: str,
) -> dict[str, Any]:
    env = envelope or {}
    route = env.get("route") or {}
    llm = env.get("llm") or {}
    card = env.get("card") or {}
    stages = env.get("stages") or []
    actions = env.get("actions") if isinstance(env.get("actions"), list) else []
    followups = env.get("followup_suggestions") if isinstance(env.get("followup_suggestions"), list) else []

    answer_sources = [
        env.get("answer_text"),
        env.get("answer"),
        env.get("message"),
        result.get("error") if isinstance(result, dict) else None,
        error,
    ]
    answer = compact(next((s for s in answer_sources if s), ""), 240)

    actual: dict[str, Any] = {
        "skill_key": env.get("skill_key") or "",
        "intent": env.get("intent") or route.get("intent") or "",
        "template_id": env.get("template_id") or card.get("templateId") or "",
        "answer": answer,
        "result_type": (result.get("result_type") if isinstance(result, dict) else "")
        or ("chat_envelope" if item["kind"] == "chat" else ""),
        "route_source": route.get("source") or "",
        "route_decision": route.get("decision") or "",
        "route_confidence": route.get("confidence", ""),
        "model_used": route.get("model_used") or llm.get("model_used") or "",
        "model_status": llm.get("model_status") or "",
        "knowledge_status": route.get("knowledge_status") or "",
        "render_ok": bool(
            env.get("rendered_html")
            or card.get("renderStatus") == "ok"
            or route.get("render_status") == "ok"
        ),
        "actions": actions,
        "followups": followups,
        "stages": stages,
        "elapsed_ms": elapsed_ms,
        "error": error,
        "conversation_id": conversation_id,
    }
    if item["kind"] == "action" and not envelope:
        actual["skill_key"] = ""
        actual["template_id"] = ""
        actual["result_type"] = (result.get("result_type") if isinstance(result, dict) else "") or ""
        err_msg = ""
        if isinstance(result, dict):
            err_msg = result.get("error") or ""
        actual["answer"] = compact(err_msg or error or json.dumps(result or {}, ensure_ascii=False), 240)
    return actual


def judge_route(item: dict[str, Any], actual: dict[str, Any]) -> str:
    if item.get("negative_expectation") == "not_nearby_resource":
        return "通过" if actual["skill_key"] != "nearby_resource" else "不通过"
    if item.get("negative_expectation") == "not_strong_travel_route":
        ok = actual["skill_key"] != "travel_route" or actual["route_decision"] != "accept"
        return "通过" if ok else "不通过"
    if not item["expected_skill"]:
        return "不通过" if actual["error"] else "人工复核"
    return "通过" if actual["skill_key"] == item["expected_skill"] else "不通过"


def judge_template(item: dict[str, Any], actual: dict[str, Any]) -> str:
    if item.get("negative_expectation") == "not_diet_card":
        return "通过" if actual["template_id"] != "diet_card" else "不通过"
    if not item["expected_template"]:
        return "人工复核" if actual["render_ok"] else "不通过"
    expected = [x.strip() for x in re.split(r"\s*或\s*|\s*/\s*", item["expected_template"]) if x.strip()]
    return "通过" if actual["template_id"] in expected else "不通过"


def judge_output(item: dict[str, Any], actual: dict[str, Any]) -> str:
    if actual["error"]:
        return "不通过"
    if not actual["answer"] or re.match(r"^抱歉，我暂时无法处理", actual["answer"]):
        if re.search(r"友好|兜底|补充|不能作确定诊断|权限不足", item["expected_output"]):
            return "通过"
        return "不通过"
    if not actual["render_ok"] and item["kind"] != "action":
        return "不通过"
    return "通过"


def judge_same_scene_actions(actual: dict[str, Any]) -> str:
    skill = actual["skill_key"]
    if not skill:
        return "不适用"
    keys = [
        *(a.get("action_key") or a.get("key") or "" for a in actual["actions"]),
        *(a.get("action_key") or a.get("key") or "" for a in actual["followups"]),
    ]
    keys = [k for k in keys if k]
    if not keys:
        return "不适用"

    if skill == "find_service":
        allowed = lambda key: key.startswith("find_service.") or key.startswith("sos.")  # noqa: E731
    else:
        allowed = lambda key: key.startswith(f"{skill}.")  # noqa: E731
    return "通过" if all(allowed(k) for k in keys) else "不通过"


def build_remark(item: dict[str, Any], actual: dict[str, Any], error: str) -> str:
    if error:
        return compact(error, 180)
    parts: list[str] = []
    if judge_route(item, actual) == "不通过":
        parts.append(f"路由偏差: expected {item['expected_skill'] or item.get('negative_expectation')}, actual {actual['skill_key']}")
    if judge_template(item, actual) == "不通过":
        parts.append(f"模板偏差: expected {item['expected_template'] or item.get('negative_expectation')}, actual {actual['template_id']}")
    if judge_output(item, actual) == "不通过":
        parts.append("输出为空、兜底或渲染异常")
    return "；".join(parts)


# ── 日志采集 ──────────────────────────────────────────────────
def collect_logs(rows: list[dict[str, Any]], client: FlatTalkClient) -> dict[str, Any]:
    ids = {r["conversation_id"] for r in rows}
    trace_resp = client.trace_logger_list(limit=5000)
    trace_items = trace_resp.get("items", []) if isinstance(trace_resp, dict) else []
    trace_rows: list[dict[str, Any]] = []
    for entry in trace_items:
        if entry.get("conversation_id") not in ids:
            continue
        route = entry.get("route") or {}
        stages = entry.get("stages") or []
        trace_rows.append({
            "ts_bj": entry.get("ts_bj", ""),
            "level": entry.get("level", ""),
            "kind": entry.get("kind", ""),
            "conversation_id": entry.get("conversation_id", ""),
            "question": compact(entry.get("question", ""), 160),
            "scene_key": route.get("scene_key", ""),
            "intent": (route.get("intent_context") or {}).get("intent_type", ""),
            "decision": route.get("decision", ""),
            "confidence": route.get("confidence", ""),
            "template": last_stage_template(stages),
            "total_ms": last_stage_ms(stages),
            "knowledge_status": route.get("knowledge_status", ""),
            "knowledge_remote_status": route.get("knowledge_remote_status", ""),
            "model_used": route.get("model_used", ""),
            "stages": " | ".join(f"{s.get('stage')}:{s.get('ms')}ms" for s in stages),
        })
    return {
        "traceRows": trace_rows,
        "runtimeTail": read_jsonl_tail(_ROOT / "data" / "runtime.log", 160),
        "eventTail": read_jsonl_tail(_ROOT / "data" / "runtime-events.log", 160),
    }


def last_stage_ms(stages: list[dict[str, Any]]) -> Any:
    return stages[-1].get("ms", "") if stages else ""


def last_stage_template(stages: list[dict[str, Any]]) -> str:
    for stage in reversed(stages):
        detail = safe_json(stage.get("detail"))
        if detail and detail.get("template_id"):
            return detail["template_id"]
    return ""


def read_jsonl_tail(file: Path, limit: int) -> list[dict[str, Any]]:
    if not file.exists():
        return []
    lines = file.read_text(encoding="utf-8").strip().splitlines()
    out: list[dict[str, Any]] = []
    for line in lines[-limit:]:
        if not line.strip():
            continue
        item = safe_json(line) or {"raw": line}
        out.append({k: (json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v) for k, v in item.items()})
    return out


# ── 模板覆盖统计 ──────────────────────────────────────────────
def build_coverage(text: str, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    coverage: list[dict[str, Any]] = []
    parts = text.split("################################################################\n# 模板覆盖清单")
    coverage_section = parts[1] if len(parts) > 1 else ""
    skill = ""
    for line in coverage_section.splitlines():
        skill_match = re.match(r"^【([^】]+)】", line)
        if skill_match:
            skill = skill_match.group(1)
            continue
        item = re.match(r"^\s*([a-zA-Z0-9_]+)[^：:]*[：:](.+)$", line)
        if not item or not skill:
            continue
        template = item.group(1)
        refs = [r for r in re.split(r"[、,，\s]+", item.group(2)) if r]
        actual_hits = sum(1 for r in rows if r["实际template_id"] == template)
        coverage.append({
            "技能": skill,
            "模板": template,
            "脚本引用": " ".join(refs),
            "本轮实际命中次数": actual_hits,
            "覆盖状态": "已覆盖" if actual_hits > 0 else "脚本未执行或未命中",
        })
    return coverage


# ── 汇总 / 评估意见 ───────────────────────────────────────────
def build_summary(
    rows: list[dict[str, Any]],
    logs: dict[str, Any],
    started_at: datetime,
    ended_at: datetime,
) -> list[dict[str, str]]:
    total = len(rows)
    route_pass = count_rows(rows, "路由是否正确", "通过")
    template_pass = count_rows(rows, "模板是否正确", "通过")
    output_pass = count_rows(rows, "输出是否合理", "通过")
    all_pass = sum(
        1 for r in rows
        if r["路由是否正确"] == "通过" and r["模板是否正确"] == "通过" and r["输出是否合理"] == "通过"
    )
    avg = round(sum(int(r.get("耗时ms") or 0) for r in rows) / total) if total else 0
    p95 = percentile([int(r.get("耗时ms") or 0) for r in rows], 0.95)
    return [
        {"指标": "run_id", "值": RUN_ID},
        {"指标": "开始时间", "值": started_at.strftime("%Y-%m-%d %H:%M:%S")},
        {"指标": "结束时间", "值": ended_at.strftime("%Y-%m-%d %H:%M:%S")},
        {"指标": "可执行用例数", "值": total},
        {"指标": "全项通过", "值": f"{all_pass}/{total}"},
        {"指标": "路由通过", "值": f"{route_pass}/{total}"},
        {"指标": "模板通过", "值": f"{template_pass}/{total}"},
        {"指标": "输出通过", "值": f"{output_pass}/{total}"},
        {"指标": "平均耗时ms", "值": avg},
        {"指标": "P95耗时ms", "值": p95},
        {"指标": "本轮后台trace条数", "值": len(logs["traceRows"])},
    ]


def build_assessment(
    rows: list[dict[str, Any]],
    logs: dict[str, Any],
    coverage: list[dict[str, Any]],
) -> list[dict[str, str]]:
    failed = [
        r for r in rows
        if r["路由是否正确"] == "不通过" or r["模板是否正确"] == "不通过" or r["输出是否合理"] == "不通过"
    ]

    by_module: list[dict[str, str]] = []
    for module, items in group_rows(rows, "模块"):
        all_pass = sum(
            1 for r in items
            if r["路由是否正确"] == "通过" and r["模板是否正确"] == "通过" and r["输出是否合理"] == "通过"
        )
        evidence = "；".join(f"{r['测试编号']}:{r['备注']}" for r in items if r["备注"])[:5]
        by_module.append({
            "类型": "模块统计",
            "对象": module or "未分类",
            "结论": f"{all_pass}/{len(items)} 全项通过",
            "证据": evidence,
            "建议": "优先修正场景词、冲突词和 intent-template-map。" if any(r["路由是否正确"] == "不通过" for r in items) else "保持回归覆盖。",
        })

    slow: list[dict[str, str]] = []
    for r in rows:
        if int(r.get("耗时ms") or 0) >= 6000:
            slow.append({
                "类型": "性能",
                "对象": r["测试编号"],
                "结论": f"耗时 {r['耗时ms']}ms",
                "证据": f"{r['实际skill_key']}/{r['实际template_id']}",
                "建议": "检查远程知识库、云诊同步、Tavily/地图富化是否需要缓存或异步化。",
            })

    uncovered: list[dict[str, str]] = []
    for c in [row for row in coverage if row["覆盖状态"] != "已覆盖"][:40]:
        uncovered.append({
            "类型": "模板覆盖",
            "对象": f"{c['技能']}/{c['模板']}",
            "结论": c["覆盖状态"],
            "证据": c["脚本引用"],
            "建议": "补充明确输入或 action 用例，避免只在覆盖清单中出现。",
        })

    failures: list[dict[str, str]] = []
    for r in failed:
        failures.append({
            "类型": "失败用例",
            "对象": r["测试编号"],
            "结论": f"{r['路由是否正确']}/{r['模板是否正确']}/{r['输出是否合理']}",
            "证据": r["备注"] or r["实际输出摘要"],
            "建议": suggest_fix(r),
        })

    log_findings = [{
        "类型": "日志",
        "对象": "query-trace.log",
        "结论": f"本轮采集 {len(logs['traceRows'])} 条后端追踪" if logs["traceRows"] else "未采集到本轮 trace",
        "证据": "；".join(f"{x['conversation_id']}:{x['scene_key']}/{x['template']}" for x in logs["traceRows"][:3]),
        "建议": "继续保留 trace 作为回归证据。" if logs["traceRows"] else "检查 trace logger 异步写入或运行路径。",
    }]

    return [*failures, *by_module, *slow, *uncovered, *log_findings]


def suggest_fix(row: dict[str, Any]) -> str:
    if row["路由是否正确"] == "不通过":
        return "补充场景 evidence/conflict 词，或调整 Supervisor 边界映射。"
    if row["模板是否正确"] == "不通过":
        return "检查 intent-template-map、模板 manifest match 和 model-service 填槽分支。"
    if row["输出是否合理"] == "不通过":
        return "补充业务数据装配或确定性填槽，避免通用抱歉兜底。"
    return "人工复核。"


# ── Excel 写入 ────────────────────────────────────────────────
def write_workbook(file: Path, *, summary, rows, logs, coverage, assessment) -> None:
    wb = Workbook()
    # 默认 sheet 需要移除，由 add_sheet 统一管理
    wb.remove(wb.active)

    add_sheet(wb, "汇总", summary)
    add_sheet(wb, "用例明细", rows)
    add_sheet(wb, "后台trace日志", logs["traceRows"])
    add_sheet(wb, "前台输入日志", [
        {
            "ts": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "conversation_id": r["conversation_id"],
            "测试编号": r["测试编号"],
            "类型": r["类型"],
            "输入或动作": r["输入或动作"],
            "action_key": r["action_key"],
        }
        for r in rows
    ])
    add_sheet(wb, "runtime日志尾部", logs["runtimeTail"])
    add_sheet(wb, "runtime-events尾部", logs["eventTail"])
    add_sheet(wb, "模板覆盖", coverage)
    add_sheet(wb, "评估意见与修改建议", assessment)
    wb.save(file)


def add_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    sheet_name = name[:31]
    ws = wb.create_sheet(title=sheet_name)
    if not rows:
        ws.append(["空"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([_to_cell_value(row.get(h, "")) for h in headers])
    # 简单自适应列宽
    for col_idx, h in enumerate(headers, start=1):
        max_len = max([len(str(h))] + [len(str(_to_cell_value(row.get(h, "")))) for row in rows[:200]])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 60)


def _to_cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


# ── Markdown 渲染 ─────────────────────────────────────────────
def render_markdown(summary: list[dict[str, str]], assessment: list[dict[str, str]]) -> str:
    def get(name: str) -> str:
        for item in summary:
            if item["指标"] == name:
                return str(item["值"])
        return ""

    failures = [item for item in assessment if item["类型"] == "失败用例"]
    lines = [
        "# flatTalk 模板脚本自动评估报告",
        "",
        f"- Run ID：{get('run_id')}",
        f"- 可执行用例数：{get('可执行用例数')}",
        f"- 全项通过：{get('全项通过')}",
        f"- 路由通过：{get('路由通过')}",
        f"- 模板通过：{get('模板通过')}",
        f"- 输出通过：{get('输出通过')}",
        f"- 平均耗时：{get('平均耗时ms')}ms，P95：{get('P95耗时ms')}ms",
        "",
        "## 主要问题",
    ]
    if failures:
        for f in failures:
            lines.append(f"- {f['对象']}：{f['结论']}；{f['证据']}；建议：{f['建议']}")
    else:
        lines.append("- 未发现自动判定失败用例。")
    lines += ["", "## 修改建议"]
    for item in [a for a in assessment if a["类型"] != "失败用例"][:20]:
        lines.append(f"- {item['类型']}/{item['对象']}：{item['结论']}。{item['建议']}")
    lines.append("")
    return "\n".join(lines)


# ── 工具函数 ──────────────────────────────────────────────────
def count_rows(rows: list[dict[str, Any]], key: str, value: str) -> int:
    return sum(1 for r in rows if r.get(key) == value)


def percentile(values: list[int], p: float) -> int:
    import math
    nums = sorted(v for v in values if isinstance(v, (int, float)))
    if not nums:
        return 0
    # 与 JS Math.ceil(len*p)-1 一致
    idx = min(len(nums) - 1, math.ceil(len(nums) * p) - 1)
    return nums[idx]


def group_rows(rows: list[dict[str, Any]], key: str) -> list[tuple[str, list[dict[str, Any]]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(row.get(key) or "", []).append(row)
    return list(grouped.items())


def append_text(left: str, right: str) -> str:
    return "；".join([s for s in [left, right] if s])


def compact(value: Any, length: int = 200) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()[:length]


def safe_json(value: Any) -> Any:
    if not value:
        return None
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(value)
    except Exception:  # noqa: BLE001
        return None


if __name__ == "__main__":
    main()
