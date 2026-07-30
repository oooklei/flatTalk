import os

def extract_docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for ti, t in enumerate(d.tables):
        for row in t.rows:
            parts.append(' | '.join(c.text.strip() for c in row.cells if c.text.strip()))
    return '\n'.join(parts)

def extract_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f'# 表：{ws.title}')
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                out.append(' | '.join(cells))
    return '\n'.join(out)

base = r'd:\GuiCare\knowledge_docs\ai膳食'
for f in ['各种食物营养成分表.xlsx', '食物营养成分_Nuwax知识库QA_20260205.xlsx']:
    p = os.path.join(base, f)
    if os.path.exists(p):
        txt = extract_xlsx(p)
        print(f'XLSX {f}: chars={len(txt)} head={txt[:80]!r}')
    else:
        print('MISSING', f)

# 找一个 docx 政策文件
for root, _, files in os.walk(r'd:\GuiCare\knowledge_docs'):
    if 'ai膳食' in root: continue
    for f in files:
        if f.lower().endswith('.docx'):
            txt = extract_docx(os.path.join(root, f))
            print(f'DOCX {f}: chars={len(txt)} head={txt[:80]!r}')
            break
    else:
        continue
    break
